import argparse
import json
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import docx
from chromadb.config import Settings
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from openai import OpenAI

BASE_DIR = Path(__file__).resolve().parent
PERSIST_DIRECTORY = BASE_DIR / "vertordb_mx"
DEFAULT_SOURCE_DIRECTORY = Path(
    os.getenv("NCC_SOURCE_DOCUMENTS_DIR", BASE_DIR / "source_documents")
)
OPENAI_EMBED_MODEL = os.getenv("OPENAI_EMBED_MODEL", "text-embedding-3-small")
CHUNK_SIZE = int(os.getenv("NCC_CHUNK_SIZE", "1500"))
CHUNK_OVERLAP = int(os.getenv("NCC_CHUNK_OVERLAP", "200"))
SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx"}


class OpenAIEmbeddingsClient(Embeddings):
    def __init__(self, model: str = OPENAI_EMBED_MODEL):
        self.model = model
        self.client = OpenAI()

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        sanitized_texts = [text.replace("\n", " ").strip() for text in texts if text.strip()]
        if not sanitized_texts:
            return []

        response = self.client.embeddings.create(
            input=sanitized_texts,
            model=self.model,
        )
        return [item.embedding for item in response.data]

    def embed_query(self, text: str) -> list[float]:
        response = self.client.embeddings.create(
            input=text.replace("\n", " ").strip(),
            model=self.model,
        )
        return response.data[0].embedding


def get_embeddings() -> OpenAIEmbeddingsClient:
    return OpenAIEmbeddingsClient()


def _read_docx(path: Path) -> str:
    document = docx.Document(path)
    return "\n".join(
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
    )


def _read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _load_document(path: Path) -> Document | None:
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        return None

    if path.suffix.lower() == ".docx":
        page_content = _read_docx(path)
    else:
        page_content = _read_text_file(path)

    page_content = page_content.strip()
    if not page_content:
        return None

    return Document(
        page_content=page_content,
        metadata={
            "source": str(path),
            "filename": path.name,
        },
    )


def load_source_documents(source_directory: Path) -> list[Document]:
    if not source_directory.exists():
        raise FileNotFoundError(
            f"Source documents directory does not exist: {source_directory}"
        )

    documents: list[Document] = []
    for path in sorted(source_directory.rglob("*")):
        if not path.is_file():
            continue
        document = _load_document(path)
        if document is not None:
            documents.append(document)

    if not documents:
        raise ValueError(f"No supported documents found in {source_directory}")

    return documents


def split_documents(documents: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )
    return splitter.split_documents(list(documents))


def _reset_persist_directory(persist_directory: Path) -> None:
    if persist_directory.exists():
        shutil.rmtree(persist_directory)
    persist_directory.mkdir(parents=True, exist_ok=True)


def persist_documents(
    source_directory: Path = DEFAULT_SOURCE_DIRECTORY,
    persist_directory: Path = PERSIST_DIRECTORY,
    reset: bool = True,
) -> dict:
    documents = load_source_documents(source_directory)
    chunks = split_documents(documents)

    if reset:
        _reset_persist_directory(persist_directory)
    else:
        persist_directory.mkdir(parents=True, exist_ok=True)

    Chroma.from_documents(
        documents=chunks,
        embedding=get_embeddings(),
        persist_directory=str(persist_directory),
        client_settings=Settings(anonymized_telemetry=False),
    )

    metadata = {
        "source_directory": str(source_directory),
        "persist_directory": str(persist_directory),
        "document_count": len(documents),
        "chunk_count": len(chunks),
        "embedding_model": OPENAI_EMBED_MODEL,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    (persist_directory / "index_metadata.json").write_text(
        json.dumps(metadata, indent=2),
        encoding="utf-8",
    )
    return metadata


def get_vectorstore(persist_directory: Path = PERSIST_DIRECTORY) -> Chroma:
    if not persist_directory.exists():
        raise FileNotFoundError(
            f"Vector store not found at {persist_directory}. Run the update command first."
        )

    return Chroma(
        persist_directory=str(persist_directory),
        embedding_function=get_embeddings(),
        client_settings=Settings(anonymized_telemetry=False),
    )


def build_retriever(search_k: int = 3):
    return get_vectorstore().as_retriever(search_kwargs={"k": search_k})


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Embed source documents with OpenAI and persist them to Chroma."
    )
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=DEFAULT_SOURCE_DIRECTORY,
        help="Directory containing .txt, .md, or .docx files to index.",
    )
    parser.add_argument(
        "--persist-dir",
        type=Path,
        default=PERSIST_DIRECTORY,
        help="Directory where the Chroma database should be stored.",
    )
    parser.add_argument(
        "--append",
        action="store_true",
        help="Append to the existing database instead of rebuilding it from scratch.",
    )
    args = parser.parse_args()

    metadata = persist_documents(
        source_directory=args.source_dir,
        persist_directory=args.persist_dir,
        reset=not args.append,
    )
    print(json.dumps(metadata, indent=2))


if __name__ == "__main__":
    main()
